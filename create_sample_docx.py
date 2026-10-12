
from pathlib import Path

from docx import Document


output_dir = Path("test_assets")
output_dir.mkdir(exist_ok=True)

output_path = output_dir / "sample.docx"

document = Document()

document.add_heading("Aura DOCX Ingestion Test", level=1)

document.add_paragraph(
    "Aura is a multimodal AI assistant that processes documents."
)

document.add_paragraph(
    "This paragraph tests regular text extraction."
)

document.add_heading("Project Information", level=2)

table = document.add_table(rows=1, cols=2)
table.style = "Table Grid"

table.rows[0].cells[0].text = "Component"
table.rows[0].cells[1].text = "Technology"

row = table.add_row()
row.cells[0].text = "Backend"
row.cells[1].text = "FastAPI"

row = table.add_row()
row.cells[0].text = "Database"
row.cells[1].text = "PostgreSQL with pgvector"

document.save(str(output_path))

print(f"Sample DOCX created: {output_path}")
