"""No model downloads: preserve every page and OCR result in bounded windows."""
import unittest
from unittest.mock import Mock
from backend.ingestion.page_batch import page_ocr_batches

class PageBatchChecks(unittest.TestCase):
    def document(self,n=5):
        pages=[]
        for i in range(n):
            p=Mock();p.get_text.return_value=f'Page {i+1}';p.get_images.return_value=[1];p.get_drawings.return_value=[];p.number=i;pages.append(p)
        d=Mock();d.__len__=Mock(return_value=n);d.load_page.side_effect=lambda i:pages[i]
        return d,pages
    def test_order_content_and_two_page_memory_bound(self):
        d,pages=self.document();calls=[]
        def recognize(images):
            calls.append(images)
            return [(p.decode(),{'exact':p.decode()},[]) for p in images]
        rows=list(page_ocr_batches(d,lambda p:str(p.number).encode(),recognize))
        self.assertEqual([r[0] for r in rows],list(range(5)))
        self.assertEqual([r[2] for r in rows],[f'Page {i+1}' for i in range(5)])
        self.assertEqual([r[4][0] for r in rows],list('01234'))
        self.assertEqual([len(c) for c in calls],[2,2,1])
    def test_plain_text_keeps_existing_no_ocr_policy(self):
        d,pages=self.document(2)
        for p in pages:p.get_images.return_value=[]
        render=Mock();recognize=Mock()
        self.assertEqual(len(list(page_ocr_batches(d,render,recognize))),2)
        render.assert_not_called();recognize.assert_not_called()
    def test_slide_mode_reads_every_slide(self):
        d,pages=self.document(2)
        for p in pages:p.get_images.return_value=[]
        recognize=Mock(return_value=[('one',{},[]),('two',{},[])])
        self.assertEqual(len(list(page_ocr_batches(d,lambda p:b'png',recognize,always_ocr=True))),2)
        recognize.assert_called_once()
    def test_bad_render_keeps_native_text_and_next_page(self):
        d,pages=self.document(2)
        def render(p):
            if p.number==0:raise ValueError('bad image')
            return b'ok'
        rows=list(page_ocr_batches(d,render,lambda images:[('second',{},[])]))
        self.assertEqual(rows[0][2],'Page 1');self.assertIsInstance(rows[0][5],ValueError)
        self.assertEqual(rows[1][4][0],'second')

class RealDocumentBatchChecks(unittest.TestCase):
    def pdf(self):
        import pymupdf
        d=pymupdf.open()
        for i in range(3):
            page=d.new_page();page.insert_text((40,60),f'Native lesson {i+1}');page.draw_rect((30,80,100,120))
        data=d.tobytes();d.close();return data
    def test_pdf_preserves_all_native_and_ocr_page_citations(self):
        from unittest.mock import patch
        from backend.ingestion.pdf import ingest_pdf
        with patch('backend.ingestion.pdf.extract_images_text',side_effect=lambda images:[('Recognized label',{},[]) for _ in images]) as ocr:
            units,warnings=ingest_pdf(self.pdf(),'test.pdf')
        self.assertEqual(warnings,[])
        self.assertEqual({u['source_location']['page_number'] for u in units},{1,2,3})
        self.assertEqual(len([u for u in units if u['extraction_status']=='extracted']),3)
        self.assertEqual(len([u for u in units if u['extraction_status']=='ocr']),3)
        self.assertEqual([len(call.args[0]) for call in ocr.call_args_list],[2,1])

    def test_ppt_retains_every_slide_with_batched_whole_slide_ocr(self):
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Inches
        from unittest.mock import patch
        from backend.ingestion.pptx import ingest_pptx
        deck=Presentation()
        for i in range(3):
            slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text_frame.text=f'Lesson {i+1}'
        output=BytesIO();deck.save(output)
        with patch('backend.ingestion.pptx.convert_presentation_to_pdf',return_value=self.pdf()), patch('backend.ingestion.pptx.extract_images_text',side_effect=lambda images:[('Recognized label',{},[]) for _ in images]), patch.dict('os.environ',{'OLLAMA_VISION_MODEL':''}):
            units,warnings=ingest_pptx(output.getvalue(),'test.pptx')
        self.assertEqual(warnings,[])
        self.assertEqual({u['source_location']['slide_number'] for u in units},{1,2,3})
        self.assertEqual(len([u for u in units if u['extraction_status']=='rendered_slide_ocr']),3)
        self.assertEqual(len([u for u in units if u['extraction_status']=='extracted']),3)

class EmbeddingBatchOrderChecks(unittest.TestCase):
    def test_length_groups_restore_original_order_and_keep_cache(self):
        import ast,time,types
        import numpy as np
        from pathlib import Path
        from collections import OrderedDict
        from threading import RLock
        tree=ast.parse(Path('backend/ai/embeddings.py').read_text())
        cls=next(n for n in tree.body if isinstance(n,ast.ClassDef) and n.name=='EmbeddingModel')
        method=next(n for n in cls.body if isinstance(n,ast.FunctionDef) and n.name=='embed_documents')
        progress=[];ns={'np':np,'time':time,'STORAGE_DIMENSIONS':1024,'report_encoded':progress.append}
        exec(compile(ast.Module(body=[method],type_ignores=[]),'embeddings.py','exec'),ns)
        values={'zzzz':1,'a':2,'bbb':3,'cc':4};calls=[]
        def encode(texts):calls.append(list(texts));return np.array([[values[t]] for t in texts])
        model=types.SimpleNamespace(_document_lock=RLock(),_lock=RLock(),_cache=OrderedDict(),batch_size=2,_encode_documents=encode,_storage_vectors=lambda v:v)
        method=types.MethodType(ns['embed_documents'],model)
        self.assertEqual(method(['zzzz','a','bbb','a','cc']).ravel().tolist(),[1,2,3,2,4])
        self.assertEqual(calls,[['a','cc'],['bbb','zzzz']]);self.assertEqual(progress,[2,4])
        self.assertEqual(method(['cc','zzzz']).ravel().tolist(),[4,1]);self.assertEqual(len(calls),2)

if __name__=='__main__':unittest.main()
