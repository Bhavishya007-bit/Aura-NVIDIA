import io

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from backend.ingestion import ingest_file


# --------------------------------------------------
# Create a test PDF
# --------------------------------------------------

def create_test_pdf() -> bytes:

    document = pymupdf.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "ATP synthase produces ATP using "
        "the proton gradient across a membrane."
    )

    page_2 = document.new_page()

    page_2.insert_text(
        (72, 72),
        "The electron transport chain transfers "
        "electrons and pumps protons."
    )

    buffer = io.BytesIO()

    document.save(buffer)

    document.close()

    return buffer.getvalue()


# --------------------------------------------------
# Create a test DOCX
# --------------------------------------------------

def create_test_docx() -> bytes:

    document = Document()

    document.add_paragraph(
        "Cell biology is the study of cells."
    )

    document.add_paragraph(
        "Mitochondria are involved in ATP production."
    )

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()


# --------------------------------------------------
# Create a test PPTX
# --------------------------------------------------

def create_test_pptx() -> bytes:

    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1),
        Inches(8),
        Inches(1)
    )

    textbox.text = (
        "ATP synthase uses a proton gradient "
        "to produce ATP."
    )

    buffer = io.BytesIO()

    presentation.save(buffer)

    return buffer.getvalue()


# ==================================================
# PDF TEST
# ==================================================

print("\n")
print("=" * 70)
print("PDF INGESTION TEST")
print("=" * 70)

pdf_data = create_test_pdf()

pdf_result = ingest_file(
    file=pdf_data,
    filename="test.pdf"
)

print("Status:", pdf_result["status"])

print(
    "Content units:",
    len(pdf_result["content_units"])
)

for unit in pdf_result["content_units"]:

    print("-" * 50)

    print("Unit ID:", unit["unit_id"])
    print("Text:", unit["text"])

    print(
        "Location:",
        unit["source_location"]
    )


assert pdf_result["status"] == "success"

assert len(
    pdf_result["content_units"]
) == 2

assert (
    pdf_result["content_units"][0]
    ["source_location"]["kind"]
    == "page"
)

assert (
    pdf_result["content_units"][0]
    ["source_location"]["page_number"]
    == 1
)

assert (
    pdf_result["content_units"][1]
    ["source_location"]["page_number"]
    == 2
)


# ==================================================
# DOCX TEST
# ==================================================

print("\n")
print("=" * 70)
print("DOCX INGESTION TEST")
print("=" * 70)

docx_data = create_test_docx()

docx_result = ingest_file(
    file=docx_data,
    filename="test.docx"
)

print("Status:", docx_result["status"])

print(
    "Content units:",
    len(docx_result["content_units"])
)

for unit in docx_result["content_units"]:

    print("-" * 50)

    print("Unit ID:", unit["unit_id"])
    print("Text:", unit["text"])

    print(
        "Location:",
        unit["source_location"]
    )


assert docx_result["status"] == "success"

assert len(
    docx_result["content_units"]
) == 2


# ==================================================
# PPTX TEST
# ==================================================

print("\n")
print("=" * 70)
print("PPTX INGESTION TEST")
print("=" * 70)

pptx_data = create_test_pptx()

pptx_result = ingest_file(
    file=pptx_data,
    filename="test.pptx"
)

print("Status:", pptx_result["status"])

print(
    "Content units:",
    len(pptx_result["content_units"])
)

for unit in pptx_result["content_units"]:

    print("-" * 50)

    print("Unit ID:", unit["unit_id"])
    print("Text:", unit["text"])

    print(
        "Location:",
        unit["source_location"]
    )


assert pptx_result["status"] == "success"

assert len(
    pptx_result["content_units"]
) == 1

assert (
    pptx_result["content_units"][0]
    ["source_location"]["kind"]
    == "slide"
)

assert (
    pptx_result["content_units"][0]
    ["source_location"]["slide_number"]
    == 1
)


# ==================================================
# UNSUPPORTED FILE TEST
# ==================================================

print("\n")
print("=" * 70)
print("UNSUPPORTED FILE TEST")
print("=" * 70)

unsupported_result = ingest_file(
    file=b"hello world",
    filename="test.txt"
)

print(
    "Status:",
    unsupported_result["status"]
)

print(
    "Warnings:",
    unsupported_result["warnings"]
)

assert (
    unsupported_result["status"]
    == "unsupported"
)


# ==================================================
# MP4 TEST
# ==================================================

print("\n")
print("=" * 70)
print("MP4 TEST")
print("=" * 70)

fake_mp4 = (
    b"\x00\x00\x00\x18"
    b"ftyp"
    b"isom"
    b"\x00\x00\x02\x00"
)

mp4_result = ingest_file(
    file=fake_mp4,
    filename="lecture.mp4"
)

print(
    "Status:",
    mp4_result["status"]
)

print(
    "Warnings:",
    mp4_result["warnings"]
)

assert (
    mp4_result["status"]
    == "not_configured"
)

assert (
    mp4_result["content_units"]
    == []
)


# ==================================================
# FINAL RESULT
# ==================================================

print("\n")
print("=" * 70)
print("INGESTION TEST: PASS")
print("=" * 70)