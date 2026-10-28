from .page_batch import page_ocr_batches
from backend.preparation_progress import report
import io
import os
import pymupdf
from pptx import Presentation
from .ocr import extract_images_text
from .page_render import render_for_ocr
from .utils import generate_unit_id, has_zip_signature
from .ppt_convert import convert_ppt_to_pptx, convert_presentation_to_pdf, PptConversionError


def _shapes(shapes):
    for shape in shapes:
        children = getattr(shape, 'shapes', None)
        if children is not None:
            yield from _shapes(children)
        else:
            yield shape


def ingest_pptx(data: bytes, filename: str) -> tuple[list[dict], list[str]]:
    if filename.lower().endswith('.ppt'):
        try:
            data = convert_ppt_to_pptx(data)
        except (PptConversionError, ValueError) as error:
            print(f"[PPT] Conversion failed: {error}", flush=True)
            return [], [str(error)]
    if not has_zip_signature(data):
        return [], ['File does not have a valid Office ZIP signature.']
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as error:
        return [], [f'Failed to open PPTX: {error}']
    units, warnings = [], []
    # Whole-slide rendering catches SmartArt, vector labels, picture fills and
    # layout context that embedded-image extraction cannot see. If unavailable,
    # native text and embedded OCR still run and the limitation is explicit.
    rendered = None
    try:
        rendered = pymupdf.open(stream=convert_presentation_to_pdf(data), filetype='pdf')
        if len(rendered) != len(presentation.slides):
            rendered.close()
            rendered = None
            raise PptConversionError('Rendered slide count does not match the original.')
    except Exception:
        warnings.append('Whole-slide reading unavailable. Native text and embedded pictures were read; some vector/SmartArt labels may be missing. Install LibreOffice or upload a PDF exported with all slides, then prepare again.')
    report("reading_slides", 0, len(presentation.slides))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts, images = [], []
        for shape in _shapes(slide.shapes):
            if getattr(shape, 'has_table', False):
                text_parts.extend(' | '.join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
            elif getattr(shape, 'has_text_frame', False):
                text = shape.text.strip()
                if text:
                    text_parts.append(text.replace('\v', '\n'))
            try:
                image = getattr(shape, 'image', None)
                if image is not None:
                    images.append((image.blob, shape))
            except Exception as error:
                warnings.append(f'Slide {slide_number}: image extraction failed ({type(error).__name__}).')
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ''
        source = {'filename': filename, 'file_type': 'pptx'}
        native = '\n\n'.join(part for part in text_parts if part.strip()).strip()
        if native:
            units.append({'unit_id': generate_unit_id(filename, 'slide', str(slide_number)),
                          'source': source, 'text': native,
                          'source_location': {'kind': 'slide', 'slide_number': slide_number},
                          'extraction_status': 'extracted', 'warnings': []})
        before = len(units)
        recognized = extract_images_text([blob for blob, _ in images])
        for image_number, ((blob, shape), recognition) in enumerate(zip(images, recognized), start=1):
            text, metadata, image_warnings = recognition
            warnings.extend(f'Slide {slide_number}, image {image_number}: {warning}' for warning in image_warnings)
            if not text:
                continue
            location = {'kind': 'slide', 'slide_number': slide_number, 'image_number': image_number,
                        'image_bbox_emu': {'left': int(shape.left), 'top': int(shape.top),
                                           'width': int(shape.width), 'height': int(shape.height)}, **metadata}
            units.append({'unit_id': generate_unit_id(filename, 'slide_image', f'{slide_number}-{image_number}'),
                          'source': source,
                          'text': f'Slide {slide_number}: {title}\nSlide text:\n{native[:450]}\nImage {image_number} recognized text (OCR; diagram relationships are not inferred):\n{text}',
                          'source_location': location, 'extraction_status': 'ocr', 'warnings': image_warnings})
        if not native and len(units) == before:
            warnings.append(f'Slide {slide_number} contains no recognized text.')
        report("reading_slides", slide_number, len(presentation.slides))
    if rendered is not None:
        try:
            for index, page, visible_text, png, recognition, reading_error in page_ocr_batches(rendered, render_for_ocr, extract_images_text, always_ocr=True):
                number = index + 1
                report('reading_slide_images', index, len(rendered))
                try:
                    if reading_error is not None:
                        raise reading_error
                    text, metadata, notices = recognition
                    warnings.extend(f'Slide {number}: {notice}' for notice in notices)
                    # Extract text rendered by LibreOffice too: some SmartArt
                    # labels become PDF text even when OCR cannot read them.
                    combined = '\n'.join(dict.fromkeys(t for t in (visible_text, text) if t))
                    if combined:
                        units.append({'unit_id': generate_unit_id(filename, 'slide_render', str(number)),
                            'source': {'filename': filename, 'file_type': 'pptx'},
                            'text': f'Slide {number} recognized text (OCR; diagram relationships are not inferred):\n{combined}',
                            'source_location': {'kind': 'slide', 'slide_number': number,
                                                'whole_slide_render': True, **metadata},
                            'extraction_status': 'rendered_slide_ocr', 'warnings': notices})
                    if os.getenv('OLLAMA_VISION_MODEL', '').strip() and (page.get_images() or page.get_drawings()):
                        # Optional, already-configured vision reads the actual
                        # slide, never guesses relationships from OCR ordering.
                        from backend.figure_tools import interpret_image
                        visual = interpret_image(png, 'Describe the visible educational content, labels and diagram relationships on this slide. Mark uncertainty; do not add textbook facts.')
                        if visual['answerable']:
                            units.append({'unit_id': generate_unit_id(filename, 'slide_vision', str(number)),
                                'source': {'filename': filename, 'file_type': 'pptx'},
                                'text': f'Model interpretation of slide {number} (not independently verified; preserve these limitations):\n' +
                                    '\n'.join(visual['observations']) + '\n' + visual['explanation'] +
                                    '\nLimitations: ' + '; '.join(visual.get('limitations', [])),
                                'source_location': {'kind': 'slide', 'slide_number': number, 'visual_interpretation': True},
                                'extraction_status': 'vision_model_checked', 'warnings': []})
                        else:
                            warnings.append(f'Slide {number}: diagram interpretation could not be verified.')
                except Exception as error:
                    warnings.append(f'Slide {number}: whole-slide reading was incomplete ({type(error).__name__}).')
                report('reading_slide_images', number, len(rendered))
        finally:
            rendered.close()
    return units, list(dict.fromkeys(warnings))

